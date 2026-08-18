from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


DOCX_PATH = "/Users/efi/Documents/Work/Faith Associates/_site-audit/docs/Faith Associates - ALL Services Proposed Replacement Content.docx"


GROUPS = [
    (
        "Training Courses",
        [
            "1 Day Mosque Management & Governance Master Class Training",
            "2 Day Mosque Management & Governance Master Class Training",
            "Safer Recruitment Training",
            "First Aid Training for Mosques and Madrassahs",
            "Risk Assessment Training for Mosques and Madrassahs",
            "Mosque Security Awareness Training",
        ],
    ),
    (
        "Mosque Services",
        [
            "Protect Duty and Martyn's Law Training",
            "Beacon Mosque Accreditation",
            "Women in Mosque Leadership",
            "Mosque Policy and Procedure Development",
            "Mosque Election Management",
            "Mosque Security Risk Assessment",
        ],
    ),
    (
        "Strategic Services",
        [
            "Strategic Campaign Planning & Execution",
            "Strategic Conference and Event Planning",
            "Media Planning",
        ],
    ),
    (
        "Safeguarding Services",
        [
            "Designated Safeguarding Lead Training",
            "International Safeguarding Training",
            "E-Safety Training",
            "Child Protection & Safeguarding Training",
        ],
    ),
]


def insert_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    if style:
        new_paragraph.style = style
    return new_paragraph


def main():
    doc = Document(DOCX_PATH)

    if any("WordPress page grouping from screenshot" in p.text for p in doc.paragraphs):
        print("Screenshot grouping already present; no changes made.")
        return

    content_heading = None
    in_first_service = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("1. Practical support for stronger mosque governance."):
            in_first_service = True
            continue
        if in_first_service and text.startswith("2. "):
            break
        if in_first_service and text == "Content sections to use on page":
            content_heading = paragraph
            break

    if content_heading is None:
        raise RuntimeError("Could not find the first Mosque Services content section.")

    anchor = content_heading
    anchor = insert_after(
        anchor,
        "WordPress page grouping from screenshot",
        "Heading 3",
    )
    anchor = insert_after(
        anchor,
        "Use these grouped service-card sections on the Mosque Services page so the replacement mirrors the old WordPress page structure.",
    )

    for group_name, items in GROUPS:
        anchor = insert_after(anchor, group_name, "Heading 3")
        for item in items:
            anchor = insert_after(anchor, item, "List Bullet")

    doc.save(DOCX_PATH)
    print(f"Updated {DOCX_PATH}")


if __name__ == "__main__":
    main()
